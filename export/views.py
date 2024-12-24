import os
import shutil

from django.http import HttpResponse
from django.shortcuts import render

from docx import Document

from constructor_vo_server import settings
from programs.models import Product, Nsi, Program, Discipline, Process
from docxtpl import DocxTemplate
from django.db.models import Q



def export_data(request, program_id):
    export_dataset(program_id)
    export_prd(program_id)
    export_prd(program_id)
    export_opd(program_id)

    return archive_and_serve(request,program_id)

def export_dataset(program_id):
    # Создание нового документа
    doc = Document('export/templates/base.docx')
    program = Program.objects.get(id=program_id)
    # Получаем все продукты
    products = Product.objects.filter(program_id=program_id)

    p = doc.add_paragraph()
    p.add_run(f"Датасет программы «{program.profile}»").bold = True
    doc.add_paragraph(f"Уровень образования: {program.level_id.name}")
    doc.add_paragraph(f"Направление: {program.direction_id.code} {program.direction_id.name}")
    doc.add_paragraph()

    # Для каждого продукта добавляем в документ его название и список Nsi
    for product in products:
        # Добавляем название продукта жирным
        p = doc.add_paragraph()
        p.add_run(product.name).bold = True
        # Добавляем описание
        doc.add_paragraph(product.description)
        # Получаем связанные Nsi и добавляем их в список
        nsi_list = product.nsis.all()
        if nsi_list:
            # Добавляем заголовок для списка Nsi
            doc.add_paragraph("Нормативные акты :")

            # Перечисляем все связанные Nsi
            for nsi in nsi_list:
                doc.add_paragraph(nsi.nsiFullName, style='marked_list')

        # Добавляем пустую строку после каждого продукта
        doc.add_paragraph()

    # Сохраняем документ
    directory = os.path.join(f"export/programs/{program_id}")
    if not os.path.exists(directory):
        os.makedirs(directory)

    url = f"export/programs/{program_id}/Датасет.docx"
    doc.save(url)



def export_rd(program_id):
    # Создание нового документа

    program = Program.objects.get(id=program_id)
    # Получаем все продукты
    products = Product.objects.filter(program_id=program_id)

    # Проверяем директорию
    directory = os.path.join(f"export/programs/{program_id}/Реконструкция деятельности")
    if not os.path.exists(directory):
        os.makedirs(directory)

    for product in products:
        doc = Document('export/templates/base.docx')

        # Заголовок
        p = doc.add_paragraph("ДЕКОМПОЗИЦИЯ ЖИЗНЕННОГО ЦИКЛА", style='bold_centered')
        run = p.add_run()
        run.add_break()
        run = p.add_run(f"продукта «{product.name}»")

        # Список и описание этапов
        p = doc.add_paragraph()
        p.add_run("1. Этапы жизненного цикла и их краткая аннотация").bold = True
        stages = product.stages.order_by('position');
        for stage in stages:
            p = doc.add_paragraph()
            p.add_run(f"1.{stage.position} {stage.name}").bold = True
            # описание этапа
            description_lines = split_to_lines(stage.description)
            for line in description_lines:
                doc.add_paragraph(line, style='main_text_12')

        # пустая строка
        p = doc.add_paragraph()

        # вторая часть
        p = doc.add_paragraph()
        p.add_run("2. Процессы жизненного цикла продукта, результаты деятельности и регулирующие их нормативные акты").bold = True
        for stage in stages:
            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Этап {stage.position} {stage}").bold=True

            p = doc.add_paragraph("Перечень процессов:",style='main_text_12')
            processes = stage.processes.all()

            for process in processes:
                p = doc.add_paragraph(f"{stage.position}.{process.position} {process.name}", style='marked_list_12')

        for stage in stages:
            p = doc.add_paragraph()
            p.add_run(f"Этап {stage.position} {stage}").bold = True
            processes = stage.processes.all()
            for process in processes:
                p = doc.add_paragraph()
                p.add_run(f"Процесс {process.position} {process.name}").bold = True
                description_lines = split_to_lines(process.description)
                for line in description_lines:
                    doc.add_paragraph(line, style='main_text_12')

                p = doc.add_paragraph(style='main_text_12')
                p.add_run(f"Результат деятельности").bold = True
                doc.add_paragraph(process.result, style='main_text_12')

                p = doc.add_paragraph(style='main_text_12')
                p.add_run(f"Нормативные акты").bold = True

                nsi_list = process.nsis.all()
                for nsi in nsi_list:
                    doc.add_paragraph(nsi.nsiFullName, style='marked_list_12')
                doc.add_paragraph()
        # СОохранение
        url = f"{directory}/{product.name}.docx"
        doc.save(url)



    return True


def export_opd(program_id):
    program = Program.objects.get(id=program_id)
    competences = program.competences.filter(type='Общепрофессиональные')

    doc = Document('export/templates/base.docx')

    # Заголовок
    p = doc.add_paragraph("ДЕКОМПОЗИЦИЯ ЖИЗНЕННОГО ЦИКЛА", style='bold_centered')
    run = p.add_run()
    doc.add_paragraph()

    for comp in competences:
        p = doc.add_paragraph()
        p.add_run(f"{comp.code} {comp.name}").bold = True
        for disc in comp.disciplines.all():
            p = doc.add_paragraph()
            p.add_run(f"Дисциплина: {disc.name}").bold = True

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Предметная область").bold = True
            doc.add_paragraph(f"{disc.area}", style='main_text_12')

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Содержание").bold = True
            description_lines = split_to_lines(disc.description)
            for line in description_lines:
                doc.add_paragraph(line, style='main_text_12')

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Практическое задание").bold = True
            doc.add_paragraph(f"{disc.task}", style='main_text_12')

    url = f"export/programs/{program_id}/ОПД.docx"
    doc.save(url)

    return True



def export_prd(program_id):
    program = Program.objects.get(id=program_id)
    # Получаем все продукты
    products = Product.objects.filter(program_id=program_id)
    directory = os.path.join(f"export/programs/{program_id}/Профессиональные дисциплины")
    if not os.path.exists(directory):
        os.makedirs(directory)

    for product in products:
        doc = Document('export/templates/base.docx')

        # Заголовок
        p = doc.add_paragraph("ПРЕОБРАЗОВАНИЕ", style='bold_centered')
        run = p.add_run()
        run.add_break()
        run = p.add_run(f"реконструкции деятельности по продукту «{product.name}» в дисциплины учебного плана")


        p = doc.add_paragraph()
        p.add_run("1. Исходные элементы дейстельности").bold = True
        stages = product.stages.order_by('position')
        for stage in stages:
            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Этап {stage.position} {stage}").bold = True

            # p = doc.add_paragraph("Перечень процессов:", style='main_text_12')
            processes = stage.processes.order_by('position')

            for process in processes:
                p = doc.add_paragraph(f"{stage.position}.{process.position} {process.name}", style='marked_list_12')

        # пустая строка
        p = doc.add_paragraph()

        # вторая часть
        p = doc.add_paragraph()
        p.add_run("2. Профессиональные дисциплины").bold = True

        disciplines = Discipline.objects.filter(
            Q(id__in=product.stages.values('discipline_id')) |  # дисциплины через Stage
            Q(id=product.discipline_id) |  # дисциплина самого продукта
            Q(id__in=Process.objects.filter(stage__in=product.stages.all()).values('discipline_id'))
            # дисциплины через Process
        ).distinct()

        for disc in disciplines:
            p = doc.add_paragraph()
            p.add_run(f"Дисциплина: {disc.name}").bold = True

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Деятельность").bold = True

            products = disc.products.order_by('position')
            stages = disc.stages.order_by('position')
            processes = disc.processes.order_by('position')

            for product in products:
                doc.add_paragraph(f"Продукт {product.name}", style="marked_list_12")

            for stage in stages:
                doc.add_paragraph(f"Этап {stage.position}. {stage.name}", style="marked_list_12")

            for process in processes:
                doc.add_paragraph(
                    f"Процесс {process.stage.position}.{process.position} {process.name}",
                    style="marked_list_12"
                )

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Содержание").bold = True
            description_lines = split_to_lines(disc.description)
            for line in description_lines:
                doc.add_paragraph(line, style='main_text_12')

            p = doc.add_paragraph(style='main_text_12')
            p.add_run(f"Практическое задание").bold = True
            doc.add_paragraph(f"{disc.task}", style='main_text_12')

        # СОохранение
        url = f"{directory}/{product.name}.docx"
        doc.save(url)

    return True


def archive_and_serve(request, program_id):
    directory = os.path.join(f"export/programs/{program_id}")
    archive_name = f'{program_id}_archive.zip'
    archive_path = os.path.join(f"export/programs/", archive_name)
    # Архивируем папку
    shutil.make_archive(archive_path.replace('.zip', ''), 'zip', directory)

    # Удаляем исходную папку
    shutil.rmtree(directory)

    # Возвращаем архив как файл для скачивания
    with open(archive_path, 'rb') as f:
        response = HttpResponse(f.read(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename={archive_name}'

    return response



def get_pdf(request):
    name = '1_Концептульное проектирование.pdf'
    path = os.path.join(f"export/files/", name)
    with open(path, 'rb') as f:
        response = HttpResponse(f.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename={name}'

    return response

def test():
    # таблица процессов
    table = doc.add_table(rows=1, cols=3)
    table.style = 'base_table'

    # Добавление заголовков столбцов
    hdr_cells = table.rows[0].cells

    start_cell(hdr_cells[0], 'Процесс и его краткая аннотация', 'table_header')
    start_cell(hdr_cells[1], 'Результат деятельности', 'table_header')
    start_cell(hdr_cells[2], 'Нормативные акты', 'table_header')

    # Для каждого процесса в stage создаём строку в таблице
    for process in stage.processes.all():
        # Добавляем строку в таблицу
        row_cells = table.add_row().cells
        # Заполняем первую колонку — процесс и его аннотация
        start_cell(row_cells[0], f"Процесс {process.position}: {process.name}", 'bold_table_content')

        description_lines = split_to_lines(process.description)
        for line in description_lines:
            p = row_cells[0].add_paragraph(line, style='table_content')

        # Заполняем вторую колонку — результат деятельности
        start_cell(row_cells[1], process.result if process.result else "Результат не указан", style='table_content')

        # Заполняем третью колонку — нормативные акты
        nsi_list = process.nsis.all()
        if nsi_list:
            nsi_text = "\n".join([nsi.nsiFullName for nsi in nsi_list])
        else:
            nsi_text = "Нормативные акты не указаны"
        start_cell(row_cells[2], nsi_text, 'table_content')


def split_to_lines(source):
    return source.split('\r\n')


def start_cell(cell, text, style):
    # Удаляем все параграфы в ячейке
    for paragraph in cell.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)

    # Добавляем новый параграф с текстом и стилем
    p = cell.add_paragraph(text)
    p.style = style
    return cell
